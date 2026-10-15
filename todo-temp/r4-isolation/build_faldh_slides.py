"""Build the FALDH failure-attribution deck (seminar 2026-10-16) on the Academic Serif B template.

Story: Round 4 wrote a principle that covers FALDH, yet FALDH stayed unannotated. Three steps
(rule-chain trace, isolation replay, single deletion) locate the blocker in three named
"do not annotate" lines written by Guideline Refinement, which sees verified examples but no gold.

Evidence (every number on the slides is read or asserted from these files):
    todo-temp/20261014_r4_failure_attribution_report.md
    todo-temp/r4-rule-chain/20261014_r4_rule_chain_report.md
    todo-temp/r4-isolation/20261014_r4_isolation_report.md
    todo-temp/r4-isolation/scores.json, scores_c3minus.json

    python todo-temp/r4-isolation/build_faldh_slides.py
"""
from __future__ import annotations

import copy
import json
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import (XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_TICK_LABEL_POSITION,
                             XL_TICK_MARK)
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = ROOT.parent / "presentations/academic-serif-template/Academic_Serif_B.pptx"
OUT = ROOT.parent / "presentations/20261016_faldh_attribution_en.pptx"  # all decks live in 修論/presentations
ISO = ROOT / "todo-temp/r4-isolation"
# Base paper Figure 1, cropped from archive/references/backbone-research.pdf page 3,
# bbox (95.6, 86.1, 499.7, 221.5) pt plus 4 pt margin, rendered at 400 dpi with PyMuPDF.
FIG1 = ISO / "assets/backbone_fig1_framework.png"
FIG1_ASPECT = 2291 / 798

# Template palette (Academic_Serif_B_Format_Guide.pptx, slide 2; accents 4 and 5 from the theme)
NAVY, INK, MUTED = "142D45", "182C3F", "52677A"
BLUE, STRONG, PALE, GRAY, BORDER = "7399C6", "315D8A", "EAF1F8", "B8C4D0", "DCE4EC"
GREEN, RED, WHITE = "336E5B", "9A4C4C", "FFFFFF"
FONT = "Arial"


# ---------------------------------------------------------------- evidence
def load_evidence() -> dict:
    scores = json.loads((ISO / "scores.json").read_text(encoding="utf-8"))
    minus = json.loads((ISO / "scores_c3minus.json").read_text(encoding="utf-8"))
    units = [u for u in scores["units"] if u["target"] == "FALDH"]
    ev = {}
    for cond in ("c1_R3", "c2_principle", "c2b_no_exception", "c3_R4"):
        rows = [u for u in units if u["condition"] == cond]
        assert len(rows) == 3, cond
        ev[cond] = {"disease": sum(u["target_status"] != "O" for u in rows),
                    "exact": sum(u["target_correct"] for u in rows),
                    "doc": sorted({(u["tp"], u["fp"], u["fn"]) for u in rows})}
    rows = [u for u in minus["units"] if u["target"] == "FALDH"]
    assert len(rows) == 3
    ev["c3_minus"] = {"disease": sum(u["target_status"] != "O" for u in rows),
                      "exact": sum(u["target_correct"] for u in rows),
                      "doc": sorted({(u["tp"], u["fp"], u["fn"]) for u in rows})}
    # The slide text states these values; fail loudly if the record ever disagrees.
    assert (ev["c1_R3"]["disease"], ev["c3_R4"]["disease"]) == (0, 0)
    assert (ev["c2_principle"]["disease"], ev["c2b_no_exception"]["disease"]) == (3, 3)
    assert ev["c2_principle"]["exact"] == ev["c2b_no_exception"]["exact"] == 0
    assert ev["c3_minus"]["exact"] == 3 and ev["c3_minus"]["doc"] == [(14, 0, 0)]
    assert ev["c3_R4"]["doc"] == [(13, 0, 1)]
    return ev


# ---------------------------------------------------------------- helpers
def rgb(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_)


def layout(prs, name):
    return next(l for l in prs.slide_layouts if l.name == name)


def add_footers(slide) -> None:
    """python-pptx skips footer and slide-number placeholders; copy them from the layout."""
    for ph in slide.slide_layout.placeholders:
        kind = ph.placeholder_format.type
        if kind not in (PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER):
            continue
        el = copy.deepcopy(ph._element)
        if kind == PP_PLACEHOLDER.SLIDE_NUMBER and not el.findall(".//" + qn("a:fld")):
            p = el.find(".//" + qn("a:p"))
            for r in p.findall(qn("a:r")):
                p.remove(r)
            fld = etree.SubElement(p, qn("a:fld"), id="{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}", type="slidenum")
            etree.SubElement(fld, qn("a:rPr"), lang="en-US")
            etree.SubElement(fld, qn("a:t")).text = "‹#›"
        slide.shapes._spTree.append(el)


def new_slide(prs, layout_name, notes):
    s = prs.slides.add_slide(layout(prs, layout_name))
    add_footers(s)
    s.notes_slide.notes_text_frame.text = notes
    return s


def place(slide, idx, x=None, y=None, w=None, h=None) -> None:
    """Move or resize a placeholder. All four values are written, because a partial
    xfrm on an inheriting placeholder leaves the other coordinates at zero."""
    ph = slide.placeholders[idx]
    lx, ly, lw, lh = ph.left, ph.top, ph.width, ph.height  # inherited from the layout
    ph.left = Inches(x) if x is not None else lx
    ph.top = Inches(y) if y is not None else ly
    ph.width = Inches(w) if w is not None else lw
    ph.height = Inches(h) if h is not None else lh


def drop(slide, idx) -> None:
    el = slide.placeholders[idx]._element
    el.getparent().remove(el)


def style_run(run, size=None, bold=None, color=None, italic=None, font=None):
    f = run.font
    if size is not None:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if color is not None:
        f.color.rgb = rgb(color)
    if font is not None:
        f.name = font


def fill_paragraphs(tf, paras, bullets=False, space_after=None, align=None, **run_defaults):
    """paras: list of str or list of (text, opts) runs. Formatting not given inherits from the layout."""
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    first = tf.paragraphs[0]
    for r in list(first.runs):
        r._r.getparent().remove(r._r)
    for i, para in enumerate(paras):
        p = first if i == 0 else tf.add_paragraph()
        runs = para if isinstance(para, list) else [(para, {})]
        for text, opts in runs:
            r = p.add_run()
            r.text = text
            style_run(r, **{**run_defaults, **opts})
        if space_after is not None:
            p.space_after = Pt(space_after)
        if align is not None:
            p.alignment = align
        if bullets:
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", "279400")
            pPr.set("indent", "-127000")
            etree.SubElement(pPr, qn("a:buChar"), char="•")


def put(slide, idx, paras, **kw):
    fill_paragraphs(slide.placeholders[idx].text_frame, paras if isinstance(paras, list) else [paras], **kw)


def text(slide, x, y, w, h, paras, size=20, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, inset=0.0, space_after=0, name=None, italic=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(inset))
    fill_paragraphs(tf, paras if isinstance(paras, list) else [paras], space_after=space_after, align=align,
                    size=size, color=color, bold=bold, font=FONT, italic=italic)
    return tb


def box(slide, x, y, w, h, fill=PALE, line=None, name=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        shp.name = name
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(fill)
    if line:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(0.75)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def boxed_text(slide, x, y, w, h, paras, fill=PALE, inset=0.2, anchor=MSO_ANCHOR.MIDDLE, name=None, **kw):
    shp = box(slide, x, y, w, h, fill=fill, name=name)
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(inset))
    kw.setdefault("align", PP_ALIGN.LEFT)
    defaults = {"size": kw.pop("size", 20), "color": kw.pop("color", INK), "bold": kw.pop("bold", False),
                "font": FONT}
    fill_paragraphs(tf, paras if isinstance(paras, list) else [paras], **kw, **defaults)
    return shp


def arrow(slide, x1, y1, x2, y2, color=STRONG):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(1.5)
    etree.SubElement(c.line._get_or_add_ln(), qn("a:tailEnd"), type="triangle", w="med", len="med")
    return c


def _border(tcPr, tag):
    ln = etree.SubElement(tcPr, qn(f"a:{tag}"), w="5715")
    sf = etree.SubElement(ln, qn("a:solidFill"))
    etree.SubElement(sf, qn("a:srgbClr"), val=BORDER)
    etree.SubElement(ln, qn("a:prstDash"), val="solid")


def style_table(gf, rows, col_w, row_h, colors=None, bold=None, center_cols=()):
    """Template table look: pale header, white body, thin #DCE4EC rules, Arial 18 pt."""
    colors, bold = colors or {}, bold or set()
    tbl = gf.table
    tblPr = tbl._tbl.tblPr
    tblPr.set("firstRow", "0")
    tblPr.set("bandRow", "0")
    for child in list(tblPr):
        tblPr.remove(child)
    for j, w in enumerate(col_w):
        tbl.columns[j].width = Inches(w)
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(row_h[i] if isinstance(row_h, (list, tuple)) else row_h)
        for j, value in enumerate(row):
            cell = tbl.cell(i, j)
            header = i == 0
            fill_paragraphs(cell.text_frame, [value], size=18, font=FONT,
                            bold=header or (i, j) in bold,
                            color=colors.get((i, j), NAVY if header else INK),
                            align=PP_ALIGN.CENTER if j in center_cols else PP_ALIGN.LEFT)
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            for child in list(tcPr):
                tcPr.remove(child)
            tcPr.set("marL", "133350")
            tcPr.set("marR", "133350")
            tcPr.set("marT", "114300")
            tcPr.set("marB", "95250")
            tcPr.set("anchor", "ctr")
            for tag in ("lnL", "lnR", "lnT", "lnB"):
                _border(tcPr, tag)
            sf = etree.SubElement(tcPr, qn("a:solidFill"))
            etree.SubElement(sf, qn("a:srgbClr"), val=PALE if header else WHITE)
    return tbl


def table_at_placeholder(slide, idx, rows, col_w, row_h, **kw):
    gf = slide.placeholders[idx].insert_table(len(rows), len(rows[0]))
    style_table(gf, rows, col_w, row_h, **kw)
    return gf


def table_at(slide, x, y, rows, col_w, row_h, **kw):
    gf = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y),
                                Inches(sum(col_w)), Inches(row_h * len(rows) if not isinstance(row_h, list) else sum(row_h)))
    style_table(gf, rows, col_w, row_h, **kw)
    return gf


def label_above(slide, txt, y=1.98):
    """Small evidence label above a chart or table, as on template slides 10 to 13."""
    return text(slide, 0.833, y, 11.667, 0.38, txt, size=18, color=MUTED)


# ---------------------------------------------------------------- slides
def build(out: Path = OUT) -> None:
    ev = load_evidence()
    prs = Presentation(TEMPLATE)
    sld_ids = prs.slides._sldIdLst
    for sld_id in list(sld_ids):  # drop the 20 sample slides
        prs.part.drop_rel(sld_id.rId)
        sld_ids.remove(sld_id)

    # 1 Cover
    s = new_slide(prs, "01 Cover", "Topic: one case from Round 4 of the NCBI dev10 GPT-5.4 run (20260802). "
                  "The Round 4 principle covers FALDH, yet FALDH stayed unannotated. "
                  "This talk traces why, in three steps.")
    put(s, 0, "The University of Tokyo")
    put(s, 1, ["Why the new principle", "did not fix FALDH"])
    put(s, 2, "Hengyu Zhou")
    put(s, 3, "Lab seminar, 16 October 2026")

    # 2 Background: the base paper's framework (its Figure 1, cropped, not redrawn)
    s = new_slide(prs, "07 Full figure",
                  "Figure 1 of the base paper (archive/references/backbone-research.pdf, page 3). Each iteration: "
                  "LLM Annotation of the dev set with the current guidelines, Evaluation against gold annotation "
                  "(strict F1), then Moderation if not satisfactory: Discrepancy Analysis groups the discrepancies "
                  "and selects the most frequent group; Pattern Explanation, Principle Generation and Guideline "
                  "Refinement (paper section 3.4) produce the next guideline. Code: infer_discrepancy_patterns, "
                  "generate_moderation_principle, refine_guidelines. The principle is an IF-THEN rule with an "
                  "explicit negative constraint, which we call the exception rule. R0 is the original guideline, "
                  "Rn the guideline after n moderation rounds. Run used here: 20260802_gpt54-high_moderation, "
                  "GPT-5.4 high effort, dev10, seed 42.")
    put(s, 0, "Background: iterative moderation framework")
    label_above(s, "Rn: guideline after n moderation rounds")
    drop(s, 1)
    fig_w = min(3.54 * FIG1_ASPECT, 11.667)
    fig = s.shapes.add_picture(str(FIG1), Inches(0.833 + (11.667 - fig_w) / 2), Inches(2.42), width=Inches(fig_w))
    fig.name = "base-paper-figure-1"
    fig._element.nvPicPr.cNvPr.set("descr", "Figure 1 of the base paper: overview of the iterative moderation "
                                   "framework, from LLM Annotation and Evaluation to Moderation (Discrepancy "
                                   "Analysis, Pattern Explanation, Principle Generation, Guideline refinement).")
    put(s, 2, "Source: base paper, Figure 1")

    # 3 The case
    s = new_slide(prs, "07 Full figure",
                  "Document PMID 10577908 (NCBI Disease train, in dev10). Gold: SpecificDisease, span [245,295), "
                  "'deficient activity of fatty aldehyde dehydrogenase'. Predictions R0 to R4: not annotated. "
                  "Round 4 targeted three SpecificDisease misses: FALDH, 'deficiency in G6PD', and "
                  "'deficiency of hepatic phenylalanine hydroxylase' (PAH).")
    put(s, 0, "The case: a disease never annotated")
    drop(s, 1)
    boxed_text(s, 0.833, 2.29, 11.667, 1.55,
               [[("Sjogren-Larsson syndrome (SLS) is an autosomal recessive disorder characterized by ichthyosis, "
                  "mental retardation, spasticity, and ", {}),
                 ("deficient activity of fatty aldehyde dehydrogenase", {"bold": True, "color": NAVY}),
                 (" (FALDH).", {})]],
               size=20, inset=0.25, name="case-sentence")
    facts = [("Gold annotation", "SpecificDisease, whole phrase", INK),
             ("Model, R0 to R4", "Not annotated in any round", RED),
             ("Round 4", "Targeted, with two other misses", INK)]
    for i, (lab, val, col) in enumerate(facts):
        y = 4.15 + i * 0.6
        text(s, 0.833, y, 3.4, 0.45, lab, size=20, color=MUTED, bold=True, name=f"fact-label-{i}")
        text(s, 4.33, y, 8.17, 0.45, val, size=24, color=col, name=f"fact-value-{i}")
    put(s, 2, "NCBI Disease, PMID 10577908; GPT-5.4, dev10")

    # 4 Research question
    s = new_slide(prs, "03 Research question",
                  "Dev F1 rose from 0.8873 (after R3) to 0.9091 (after R4), but per entity Round 4 fixed none of the "
                  "two long-standing targets; FALDH and G6PD stayed missed, PAH only undid an R3 regression. "
                  "This deck covers FALDH only.")
    put(s, 0, "Research question")
    put(s, 1, ["Why did a matching principle", "not fix FALDH?"])
    put(s, 2, "Dev F1 0.8873 → 0.9091; FALDH missed")

    # 5 Two opposite instructions
    s = new_slide(prs, "14 Case comparison",
                  "Left: the R4 principle as merged into guideline lines 12 to 14. Right: three lines in the same R4 "
                  "guideline that name FALDH as not annotated; we call them the named exclusions. "
                  "Line 54 (block 54-58): example 'Deficient activity of fatty aldehyde dehydrogenase' marked "
                  "'Do not annotate'. Line 167: 'This also excludes isolated molecular-function or "
                  "laboratory-abnormality phrases such as deficient activity of fatty aldehyde dehydrogenase ...'. "
                  "Line 263: 'the phrase is not annotated because here it refers to enzyme activity/biochemical "
                  "abnormality'. Lines 54 and 263 descend from R3 lines 41 and 241; line 167 is new in R4.")
    put(s, 0, "One guideline, two opposite instructions")
    put(s, 1, "R4 principle: annotate", color=GREEN)
    put(s, 2, ["Deficiency + specific molecule", "In definition or causal context",
               "→ annotate as SpecificDisease", "Exception rule: not mere lab findings"], space_after=6)
    put(s, 3, "Three named exclusions", color=RED)
    put(s, 4, ["L54: example “Do not annotate”", "L167: lists FALDH as excluded",
               "L263: “the phrase is not annotated”"], space_after=6)
    place(s, 2, y=3.0, h=2.4)
    place(s, 4, y=3.0, h=2.4)
    put(s, 5, "L54, L263 from R3; L167 new")

    # 6 Approach: one question per step
    s = new_slide(prs, "09 Workflow",
                  "Step 1 reads the saved prompts and guideline diffs; no model calls. It locates where FALDH is lost "
                  "and where the named exclusions come from, but cannot show that they cause the failure. Step 2 "
                  "re-annotates the FALDH document with the principle alone (and without its exception rule). Step 3 "
                  "re-annotates with the R4 guideline minus only the named exclusions. Same model, prompt template and "
                  "output settings as the original run; gold is used only for scoring.")
    put(s, 0, "Approach: three steps, three questions")
    put(s, 1, "1  Trace")
    put(s, 2, "2  Isolate")
    put(s, 3, "3  Delete")
    steps = [("Where was FALDH lost?", "Saved prompts; no new calls"),
             ("Can the principle alone fix FALDH?", "Swap only the guideline text"),
             ("Do named exclusions block FALDH?", "Delete only those three lines")]
    for i, (question, how) in enumerate(steps):
        x = 0.833 + i * 4.167
        text(s, x, 4.0, 3.333, 0.8, question, size=20, color=INK, bold=True, align=PP_ALIGN.CENTER,
             name=f"step-question-{i}")
        text(s, x, 4.85, 3.333, 0.7, how, size=18, color=MUTED, align=PP_ALIGN.CENTER, name=f"step-how-{i}")
    for x in (4.233, 8.4):
        arrow(s, x, 3.445, x + 0.6, 3.445)
    put(s, 4, "Annotator never sees gold; model fixed")
    place(s, 4, y=5.75, h=0.5)

    # 7 Step 1: lost at Guideline Refinement
    s = new_slide(prs, "13 Ablation table",
                  "Rule chain from the saved Round 4 prompts (rounds/iteration_04), stage names as in the base paper. "
                  "Discrepancy Analysis puts FALDH in the target group 'Gold: SpecificDisease vs LLM: O (Missed)', 3 "
                  "cases. Pattern Explanation sees FALDH as an error example with its gold label and suggests "
                  "SpecificDisease in 'characterized by' contexts. Principle Generation writes a principle whose "
                  "condition covers FALDH. Guideline Refinement keeps the two R3 exclusions and adds a third (line 167). "
                  "LLM Annotation with the R4 guideline: FALDH still not annotated. The next slides show what "
                  "Guideline Refinement receives and why it writes these lines.")
    put(s, 0, "Step 1: lost at Guideline Refinement")
    label_above(s, "FALDH's rule chain in Round 4")
    table_at_placeholder(s, 2, [
        ["Stage (base paper)", "Sees gold?", "What happened to FALDH"],
        ["Discrepancy Analysis", "Yes", "Put in the target group"],
        ["Pattern Explanation", "Yes", "Suggested SpecificDisease"],
        ["Principle Generation", "Yes", "Principle covers FALDH"],
        ["Guideline Refinement", "No", "Kept two, added a third exclusion"],
        ["LLM Annotation", "No, by design", "Still not annotated"],
    ], [3.6, 2.4, 5.667], 0.55, center_cols=(1,),
        colors={(4, 0): RED, (4, 1): RED, (4, 2): RED}, bold={(4, 0), (4, 1), (4, 2)})
    put(s, 1, "Source: saved R4 prompts, guideline diff")

    # 8 Step 1: what Guideline Refinement receives
    s = new_slide(prs, "08 Figure comparison",
                  "The saved R4 refine_guidelines prompt has, in order: GOAL, CURRENT GUIDELINES (the R3 guideline), "
                  "NEW MODERATION PRINCIPLES (the R4 principle), 'CONSTRAINT: DO NOT BREAK THESE EXAMPLES' (the verified "
                  "examples) and INSTRUCTIONS. It contains no discrepancy cases and no gold. The base paper (section "
                  "3.4.4) describes verified examples as in-prompt checks so that known-correct cases are not flipped; "
                  "it also says this step gets the discrepancy context, which this implementation does not pass. Each "
                  "verified example is FILE, the full TEXT of the document, and MATCHED ANNOTATIONS: at most 5 of its "
                  "strict true positives. Gold entities the model missed are not marked anywhere.")
    put(s, 0, "Step 1: what Guideline Refinement receives")
    put(s, 1, "Inputs, in prompt order")
    put(s, 3, "Why examples are included")
    drop(s, 4)
    drop(s, 2)
    boxed_text(s, 0.833, 2.92, 5.583, 2.95,
               ["1  Current guideline (R3)", "2  New principle (R4)", "3  Verified examples (5 documents)",
                "4  Instructions",
                [("Not included: discrepancies, gold", {"bold": True, "color": RED})]],
               size=20, inset=0.25, anchor=MSO_ANCHOR.TOP, space_after=10, name="refinement-inputs")
    boxed_text(s, 6.917, 2.92, 5.583, 2.95,
               [[("Prompt text:", {"bold": True, "color": MUTED, "size": 18})],
                [("“The following examples are currently handled CORRECTLY. Your new guidelines MUST NOT cause "
                  "these to be annotated incorrectly.”", {"italic": True})],
                [("Base paper §3.4.4: guard against regression", {"color": MUTED, "size": 18})]],
               size=20, inset=0.25, anchor=MSO_ANCHOR.TOP, space_after=10, name="prompt-quote-verified")
    put(s, 5, "Next: what one example looks like")

    # 8b Step 1: one verified example, as it appears in the prompt
    s = new_slide(prs, "07 Full figure",
                  "Excerpt of the verified-example block for document 10577908 in the saved Round 4 Guideline Refinement "
                  "prompt (rounds/iteration_04/refine_guidelines.txt), shortened only where marked with an ellipsis. "
                  "TEXT is the whole abstract as plain text; the corpus tokenization (spaces around punctuation) is kept. "
                  "MATCHED ANNOTATIONS lists at most 5 of the document's strict true positives; after R3 the document "
                  "has 13 of 14 gold entities correct, so 8 correct ones (for example ichthyosis, spasticity) are not "
                  "listed either. The missed gold entity FALDH is not marked anywhere. The red highlight is ours, for "
                  "the audience; the prompt has no highlighting.")
    put(s, 0, "Step 1: one verified example, as shown")
    label_above(s, "Excerpt, saved Round 4 prompt")
    drop(s, 1)
    boxed_text(s, 0.833, 2.42, 7.35, 3.85,
               [[("FILE: ", {"bold": True}), ("10577908.json", {})],
                [("TEXT: ", {"bold": True}),
                 ("… Sjogren-Larsson syndrome ( SLS ) is an autosomal recessive disorder characterized by ichthyosis , "
                  "mental retardation , spasticity , and ", {}),
                 ("deficient activity of fatty aldehyde dehydrogenase", {"color": RED, "bold": True}),
                 (" ( FALDH ) . …", {})],
                [("MATCHED ANNOTATIONS:", {"bold": True})],
                '- SpecificDisease [136, 139] "SLS"',
                '- DiseaseClass [207, 225] "mental retardation"',
                '- DiseaseClass [148, 176] "autosomal recessive disorder"',
                '- Modifier [1481, 1484] "SLS"',
                '- SpecificDisease [23, 47] "Sjogren-Larsson syndrome"'],
               size=18, inset=0.15, anchor=MSO_ANCHOR.TOP, space_after=0, name="prompt-quote-example-block")
    text(s, 8.5, 2.42, 4.0, 3.5,
         [[("TEXT: whole abstract", {"bold": True, "color": NAVY})],
          [("No labels inside the text", {"color": MUTED, "size": 18})],
          [("MATCHED: 5 of 13 correct", {"bold": True, "color": NAVY})],
          [("At most 5 listed per document", {"color": MUTED, "size": 18})],
          [("FALDH: in TEXT only", {"bold": True, "color": RED})],
          [("Missed gold is never marked", {"color": MUTED, "size": 18})],
          [("Red highlight is ours", {"italic": True, "color": MUTED, "size": 18})]],
         size=20, space_after=4, name="example-callouts")
    drop(s, 2)

    # 9 Step 1: why FALDH's document is shown
    s = new_slide(prs, "07 Full figure",
                  "Selection (iterative.py, _build_true_positive_examples): walk the dev documents in order and take "
                  "the first 5 that have at least one strict true positive; list at most 5 of its true positives. The "
                  "dev10 sample is drawn with seed 42 and then sorted by file name (sampling.py), so the order is by "
                  "PMID string, not by content or errors. The first five documents all have a true positive, so they "
                  "are shown in every round. FALDH's document 10577908 is fourth; after R3 it has 13 of 14 gold "
                  "entities correct, of which 5 are listed (SLS, mental retardation, autosomal recessive disorder, SLS, "
                  "Sjogren-Larsson syndrome). The documents of G6PD (6th) and PAH (8th) are never shown, and the "
                  "guideline never names them. norrin's document is second.")
    put(s, 0, "Step 1: why FALDH's document is shown")
    label_above(s, "Dev10 documents, sorted by file name")
    drop(s, 1)
    gap = 0.12
    w = (11.667 - 9 * gap) / 10
    tags = {2: ("norrin", MUTED), 4: ("FALDH", RED), 6: ("G6PD", MUTED), 8: ("PAH", MUTED)}
    text(s, 0.833, 2.42, 5 * w + 4 * gap, 0.38, "Shown as verified examples", size=18, color=NAVY, bold=True,
         align=PP_ALIGN.CENTER, name="shown-label")
    text(s, 0.833 + 5 * (w + gap), 2.42, 5 * w + 4 * gap, 0.38, "Not shown", size=18, color=MUTED,
         align=PP_ALIGN.CENTER, name="not-shown-label")
    for i in range(10):
        x = 0.833 + i * (w + gap)
        shown = i < 5
        shp = boxed_text(s, x, 2.85, w, 0.6, f"#{i + 1}", fill=PALE if shown else WHITE, inset=0.05,
                         size=20, bold=True, color=(RED if i == 3 else NAVY) if shown else MUTED,
                         align=PP_ALIGN.CENTER, name=f"doc-{i + 1}")
        if not shown:
            shp.line.color.rgb = rgb(GRAY)
            shp.line.width = Pt(0.75)
        if i + 1 in tags:
            label, col = tags[i + 1]
            text(s, x, 3.5, w, 0.38, label, size=18, color=col, bold=i == 3, align=PP_ALIGN.CENTER,
                 name=f"doc-tag-{i + 1}")
    text(s, 0.833, 4.1, 5.6, 1.9, ["Walk documents in file-name order", "First five with a correct entity",
                                    "Content and errors play no role", "Hence the same five every round"],
         size=20, space_after=6, name="selection-rule")
    boxed_text(s, 6.917, 4.1, 5.583, 1.85,
               [[("Consequence", {"bold": True, "color": MUTED, "size": 18})],
                [("FALDH (#4): shown every round", {"bold": True, "color": RED})],
                "G6PD, PAH: never shown, never named"],
               size=20, inset=0.2, anchor=MSO_ANCHOR.TOP, space_after=6, name="selection-consequence")
    put(s, 2, "Source: saved prompts, selection code")

    # 10 Step 1: how FALDH became an exclusion
    s = new_slide(prs, "07 Full figure",
                  "Inference from the saved prompts; the model's reasoning is not observed. (1) The new principle's "
                  "condition matches FALDH, so following it would annotate FALDH. (2) The verified example shows "
                  "FALDH's sentence under 'handled CORRECTLY' with FALDH unannotated and unmarked. (3) INSTRUCTIONS "
                  "item 3: 'Ensure your changes strictly support the New Principles WITHOUT invalidating the CONSTRAINT "
                  "examples. If a trade-off is required, prioritize specific sub-cases to avoid regression.' (4) The "
                  "output files FALDH under the principle's exception rule. R3: the R3 principle excludes 'a "
                  "molecular/substance entity or level'; R3 lines 41 and 241 say FALDH 'refers to enzyme activity'. R4: "
                  "the R4 exception rule excludes 'an isolated biochemical finding, mechanistic description, or "
                  "laboratory abnormality'; new line 167 lists FALDH as such a phrase. The R2 guideline does not mention "
                  "FALDH.")
    put(s, 0, "Step 1: how FALDH became an exclusion")
    label_above(s, "Inferred from prompts; reasoning not observed")
    drop(s, 1)
    flow = [("New principle", "Would annotate FALDH", "Principle's condition matches FALDH"),
            ("Verified example", "Shows FALDH unannotated", "Marked “handled CORRECTLY”"),
            ("Instruction 3", "Avoid regression on examples", "“prioritize specific sub-cases to avoid regression”"),
            ("Output", "Files FALDH under exception rule", "Reason echoes the exception rule")]
    w = 2.6
    gap = (11.667 - 4 * w) / 3
    for i, (head, claim, evidence) in enumerate(flow):
        x = 0.833 + i * (w + gap)
        text(s, x, 2.42, w, 0.36, head, size=18, color=MUTED, name=f"flow-head-{i}")
        boxed_text(s, x, 2.82, w, 1.4, claim, size=20, bold=True, inset=0.15, align=PP_ALIGN.CENTER,
                   color=RED if i == 3 else INK, name=f"flow-box-{i}")
        text(s, x, 4.35, w, 1.2, evidence, size=18, color=MUTED, name=f"flow-evidence-{i}")
        if i < 3:
            arrow(s, x + w + 0.05, 3.52, x + w + gap - 0.05, 3.52)
    put(s, 2, "Happened in R3, repeated in R4")

    # 11 Step 1: the same source in three runs
    s = new_slide(prs, "12 Results table",
                  "Same trace on two other runs (todo-temp/r4-rule-chain/20261014_r4_rule_chain_report.md, section 3). "
                  "dev10 top-1 is the run in this talk; dev10 top-4 shows the same five documents; in dev30 FALDH's "
                  "document is in the dev set but not among the five shown, and the guideline never names FALDH. "
                  "Across 5 cases in 3 runs, being shown was necessary but not sufficient for a named exclusion "
                  "(G6PD was shown in top-4 but not named). In dev30 FALDH is still missed without any named "
                  "exclusion, so the original miss has another cause, consistent with R0 to R2. Observational, not "
                  "a controlled test.")
    put(s, 0, "Step 1: the same source in three runs")
    label_above(s, "FALDH across three refinement runs")
    table_at_placeholder(s, 2, [
        ["Run", "Document shown?", "FALDH named?"],
        ["dev10, top-1 (this talk)", "Yes", "Yes"],
        ["dev10, top-4", "Yes", "Yes"],
        ["dev30", "No", "No"],
    ], [5.0, 3.333, 3.334], 0.72, center_cols=(1, 2))
    put(s, 1, "dev30: FALDH never named, still missed")

    # 12 Is the conflict the cause?
    s = new_slide(prs, "12 Results table",
                  "Step 1 shows that the conflict exists and where it comes from. It does not show that it causes the "
                  "failure: FALDH was already missed in R0 to R2, before any exclusion. Three explanations, each with "
                  "its own test. The first two predict opposite results when the principle is given alone (Step 2). "
                  "The exception rule is the principle's trailing 'DO NOT apply ...' clause.")
    put(s, 0, "Is the conflict the cause?")
    label_above(s, "Three explanations, each with a test")
    table_at_placeholder(s, 2, [
        ["Explanation", "If true, we should see", "Tested in"],
        ["Named exclusions block FALDH", "Principle alone works; R4 fails", "Steps 2, 3"],
        ["Model never reads FALDH as disease", "Fails even with principle alone", "Step 2"],
        ["Exception rule blocks FALDH", "Works without the exception rule", "Step 2"],
    ], [4.9, 4.6, 2.167], 0.72, center_cols=(2,), bold={(1, 0), (2, 0), (3, 0)})
    put(s, 1, "First two predict opposite results")

    # 13 Step 2 design
    s = new_slide(prs, "13 Ablation table",
                  "Isolation replay, 2026-10-14: 3 target documents x 4 guideline conditions x 3 repetitions = 36 "
                  "calls, GPT-5.4 high, same prompt template and output configuration as the original run. 'Label "
                  "definitions' are the four category definition sentences of the original guideline (lines 6-8, 20, "
                  "28). 'Without exception rule' removes only the principle's trailing 'DO NOT apply ...' clause. This "
                  "slide reports the FALDH document only.")
    put(s, 0, "Step 2: give the principle alone")
    label_above(s, "FALDH document, 3 runs per condition")
    table_at_placeholder(s, 2, [
        ["Guideline given to the annotator", "Question answered"],
        ["R3 guideline (before Round 4)", "Is the FALDH miss reproduced?"],
        ["Label definitions + R4 principle", "Can the principle alone fix FALDH?"],
        ["Same, without exception rule", "Does the exception rule block FALDH?"],
        ["R4 guideline (after Round 4)", "Does the full guideline block FALDH?"],
    ], [6.0, 5.667], 0.62, bold={(1, 0), (2, 0), (3, 0), (4, 0)})
    put(s, 1, "Subset of a 36-call isolation run")

    # 14 Step 2 results (chart)
    s = new_slide(prs, "10 Bar chart",
                  "Counts out of 3 runs per condition, document 10577908. 'Labeled as disease' counts any prediction "
                  "overlapping the gold span with a disease label. Principle only: SpecificDisease [245,305) twice "
                  "(span includes '( FALDH )'), CompositeMention [194,295) once; without the exception rule: the same "
                  "pattern. The R4 guideline reproduces the saved R4 output (not annotated) in all three runs. Exact "
                  "match requires the gold span and label; the principle-only conditions lack the full guideline's span "
                  "rules, so their strict score is 0/3. Source: todo-temp/r4-isolation/scores.json.")
    put(s, 0, "Step 2: the principle works on its own")
    label_above(s, "FALDH outcome, runs out of 3")
    data = CategoryChartData()
    data.categories = ["R3 guideline", "Principle only", "No exception rule", "R4 guideline"]
    order = ("c1_R3", "c2_principle", "c2b_no_exception", "c3_R4")
    data.add_series("Labeled as disease", [ev[c]["disease"] for c in order])
    data.add_series("Exact gold match", [ev[c]["exact"] for c in order])
    chart = s.placeholders[3].insert_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, data).chart
    style_chart(chart)
    put(s, 1, ["Works alone,", "fails in R4"])
    put(s, 2, "Spans add “(FALDH)”: strict 0/3")
    place(s, 2, h=0.95)

    # 15 Step 3: delete only the named exclusions
    r4, minus = ev["c3_R4"], ev["c3_minus"]
    (tp4, fp4, fn4), (tpm, fpm, fnm) = r4["doc"][0], minus["doc"][0]
    s = new_slide(prs, "08 Figure comparison",
                  "R4 guideline with only the three named exclusions deleted: lines 54-58 (the example block and its "
                  "trailing blank line), 167 and 263; everything else verbatim, including the principle in lines 12-14. "
                  "3 calls, FALDH document only. The document has 14 gold entities. With the full R4 guideline the "
                  f"other 13 are all correct ({tp4} TP, {fp4} FP, {fn4} FN: the one FN is FALDH, predicted as nothing, "
                  "so no boundary error). Without the named exclusions FALDH is exact SpecificDisease [245,295) in 3 of "
                  f"3 runs and the document is {tpm} TP, {fpm} FP, {fnm} FN in every run. "
                  "Source: todo-temp/r4-isolation/scores.json, scores_c3minus.json.")
    put(s, 0, "Step 3: delete only the named exclusions")
    put(s, 1, "R4 guideline")
    put(s, 3, "R4 without named exclusions")
    drop(s, 5)
    drop(s, 4)
    drop(s, 2)
    panels = ((0.833, r4, RED, f"{tp4} of {tp4 + fn4} correct; FALDH missed", "r4"),
              (6.917, minus, GREEN, f"{tpm} of {tpm + fnm} correct; {fpm} FP", "r4-minus"))
    for x, res, col, doc_line, nm in panels:
        box(s, x, 2.92, 5.583, 1.85, name=f"result-frame-{nm}")
        text(s, x, 2.98, 5.583, 1.0, f"{res['exact']} / 3", size=60, color=col, bold=True, align=PP_ALIGN.CENTER,
             name=f"stat-{nm}")
        text(s, x, 3.95, 5.583, 0.42, "runs with FALDH correct", size=20, color=INK, align=PP_ALIGN.CENTER,
             name=f"stat-label-{nm}")
        text(s, x, 4.38, 5.583, 0.4, doc_line, size=18, color=MUTED, align=PP_ALIGN.CENTER, name=f"doc-score-{nm}")
    text(s, 0.833, 5.05, 11.667, 1.3,
         [[("Deleted: the three named exclusions", {"bold": True, "color": MUTED})],
          "L54–58: FALDH example, “Do not annotate”",
          "L167: “also excludes … fatty aldehyde dehydrogenase”",
          "L263: “the phrase is not annotated”"],
         size=18, space_after=2, name="prompt-quote-deleted-lines")

    # 16 Verdict
    s = new_slide(prs, "12 Results table",
                  "Named exclusions: supported; deleting only them restores FALDH exactly (3/3) without new errors in "
                  "the document. Model judgment: rejected for Round 4; given the principle alone, the model labels "
                  "FALDH as a disease in 3 of 3 runs. Exception rule: rejected; removing it changes nothing. Caveat: "
                  "FALDH was already missed in R0 to R2, before any exclusion existed; that original miss is the "
                  "model's own judgment. The exclusions explain why Round 4 could not fix it. This corrects the 10/07 "
                  "statement that the exclusions were a symptom, not a cause.")
    put(s, 0, "Verdict: the conflict blocked the fix")
    label_above(s, "Checked against Steps 2 and 3")
    table_at_placeholder(s, 2, [
        ["Explanation", "Verdict", "Deciding evidence"],
        ["Named exclusions block FALDH", "Supported", "Deletion restores FALDH, 3/3"],
        ["Model never reads FALDH as disease", "Rejected", "Principle alone labels FALDH, 3/3"],
        ["Exception rule blocks FALDH", "Rejected", "Unchanged without the exception rule"],
    ], [4.9, 1.75, 5.017], 0.72, center_cols=(1,),
        colors={(1, 1): GREEN, (2, 1): RED, (3, 1): RED}, bold={(1, 0), (2, 0), (3, 0), (1, 1), (2, 1), (3, 1)})
    put(s, 1, "Original R0–R2 miss predates any exclusion")

    # 17 Conclusion
    s = new_slide(prs, "17 Summary",
                  "Evidence chain: (1) Guideline Refinement sees no gold, only verified-example documents chosen by "
                  "file name; (2) FALDH appears there unannotated and unmarked, and the step, told to avoid regression "
                  "on those examples, wrote FALDH as a named exclusion (inferred from the prompt); (3) the named "
                  "exclusions blocked a principle that works on its own; deleting them restores FALDH. Next step "
                  "(endorsed at the 10/09 seminar): give Guideline Refinement the gold annotations of the target "
                  "errors. A second idea from that seminar, checking the refined guideline against dev gold, is not "
                  "covered in this deck. This result also revises the 10/07 statement that the exclusions were only a "
                  "symptom.")
    put(s, 0, "Conclusion and next steps")
    # Bullets 2-4 as edited by the user in ONLYOFFICE (2026-10-15 23:17); keep verbatim.
    put(s, 1, ["FALDH was lost at Guideline Refinement", "That step sees full text of Verified Examples",
               "FALDH appeared in that text with no gold label",
               "The step then wrote named exclusions of FALDH (Find which rule lead to this exclusion?)",
               "Deleting the exclusions restores FALDH"], bullets=True, space_after=8)
    place(s, 1, h=2.95)
    put(s, 2, "Next: show gold to Guideline Refinement")
    place(s, 2, y=5.6)

    # 18 Limitations
    s = new_slide(prs, "14 Case comparison",
                  "Scope: each condition has 3 runs, so rates are coarse; the isolation replay covers 3 dev documents "
                  "from one dev10 run; one model and setting (GPT-5.4, high reasoning effort); the 'labeled as disease' "
                  "reading was fixed after the first round of results. Not yet tested: deleting the norrin named "
                  "exclusion; the refinement mechanism on slide 10, which is inferred from prompts; whether removing "
                  "named exclusions changes validation F1.")
    put(s, 0, "Limitations")
    put(s, 1, "Scope of the evidence")
    put(s, 2, ["Three runs per condition", "Three dev documents, one run", "Single model: GPT-5.4, high effort",
               "Criterion set after first results"], space_after=6)
    put(s, 3, "Not yet tested")
    put(s, 4, ["Deleting the norrin exclusion", "The inferred refinement mechanism", "Effect on validation F1"],
        space_after=6)
    place(s, 2, y=3.0, h=2.6)
    place(s, 4, y=3.0, h=2.6)
    drop(s, 5)

    # 19 Questions
    s = prs.slides.add_slide(layout(prs, "18 Questions"))
    put(s, 0, "Questions")
    put(s, 1, "Hengyu Zhou")
    put(s, 2, "The University of Tokyo")

    normalize_slide_ids(prs)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    print(f"saved {out} ({len(prs.slides)} slides)")


R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def normalize_slide_ids(prs) -> None:
    """Write every <p:sldId> as id="..." r:id="..." with the conventional r prefix.

    The template's presentation.xml declares the relationships namespace only on each
    sldId, not on the root, so new entries came out as ns0:id. That is valid XML and
    PowerPoint reads it, but ONLYOFFICE matches the literal r:id and showed an empty deck."""
    lst = prs.slides._sldIdLst
    for old in list(lst):
        new = etree.Element(qn("p:sldId"), nsmap={"r": R_NS})
        new.set("id", old.get("id"))
        new.set(f"{{{R_NS}}}id", old.get(f"{{{R_NS}}}id"))
        lst.replace(old, new)


def style_chart(chart) -> None:
    """Match template chart 1: gray baseline and blue highlight series, Arial 18 pt, light grid."""
    chart.font.name = FONT
    chart.font.size = Pt(18)
    chart.font.color.rgb = rgb(INK)
    plot = chart.plots[0]
    plot.gap_width = 120
    plot.overlap = 0
    for series, (fill, line) in zip(plot.series, ((GRAY, MUTED), (BLUE, STRONG))):
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = rgb(fill)
        series.format.line.color.rgb = rgb(line)
        series.format.line.width = Pt(0.75)
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format = "0"
    dl.number_format_is_linked = False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size = Pt(18)
    dl.font.color.rgb = rgb(INK)
    va = chart.value_axis
    # Counts are printed on the bars, so the value axis is hidden. Its maximum of 4 only leaves
    # room for the label above a full bar; ONLYOFFICE rounds a 3.5 maximum up to a visible 4,
    # and inside-end labels of zero bars fall below the baseline there.
    va.minimum_scale, va.maximum_scale, va.major_unit = 0, 4, 1
    va.has_major_gridlines = False
    va.visible = False
    # python-pptx writes a bare <c:delete/>, whose schema default is true; ONLYOFFICE reads it as
    # false, so write the value explicitly and also blank the labels, ticks and line.
    va._element.find(qn("c:delete")).set("val", "1")
    va.tick_label_position = XL_TICK_LABEL_POSITION.NONE
    va.major_tick_mark = XL_TICK_MARK.NONE
    va.format.line.fill.background()
    ca = chart.category_axis
    ca.format.line.fill.background()
    ca.tick_labels.font.size = Pt(18)
    ca.tick_labels.font.color.rgb = rgb(INK)
    ca.major_tick_mark = XL_TICK_MARK.NONE
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    chart.legend.font.size = Pt(18)
    chart.legend.font.color.rgb = rgb(INK)


if __name__ == "__main__":
    import sys
    build(Path(sys.argv[1]) if len(sys.argv) > 1 else OUT)  # optional output path, e.g. while the deck is open
